# =============================================================================
# tests/test_pipeline.py
#
# End-to-end test for the Document Migration Tool v3 pipeline.
#
# Creates two minimal sample .docx files (a source document and a destination
# template), then runs the full pipeline:
#
#   parse → auto-map → migrate (append mode) → migrate (replace mode)
#   → boilerplate detection → report export
#
# Run from the project root:
#   python tests/test_pipeline.py
#
# Expected result: "ALL TESTS PASSED" with 0 errors.
# =============================================================================

import sys
from pathlib import Path

# Add project root to path so all packages resolve correctly
sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document

from core.audit_log            import MigrationReportExporter
from core.auto_mapper          import AutoMapper
from core.boilerplate_detector import BoilerplateDetector
from core.content_migrator     import ContentMigrator, MigrationReport
from core.document_parser      import DocumentParser
from models.sections           import MappingStatus, MigrationMode, MatchMethod


# =============================================================================
# Sample document builders
# =============================================================================

def create_source_document(output_path: str):
    """
    Create a minimal source document (SEP v1) with:
      - Multiple heading levels to test hierarchy handling
      - One blank heading to test the empty-section filter
      - A table in one section to test table migration
      - Paragraph content under every section
    """
    doc = Document()

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("This document provides an introduction to the program.")

    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph("The purpose of this SEP is to define software engineering activities.")

    doc.add_heading("1.2 Scope", level=2)
    doc.add_paragraph("This SEP applies to all software components.")

    # Intentional blank heading — should be filtered out by DocumentParser
    doc.add_heading("", level=2)

    doc.add_heading("1.3 Background", level=2)
    doc.add_paragraph("The program was initiated in 2020.")

    doc.add_heading("2. Management Overview", level=1)
    doc.add_paragraph("This section describes the management structure.")

    doc.add_heading("2.1 Organizational Structure", level=2)
    doc.add_paragraph("The program is organized into three functional groups.")

    doc.add_heading("2.2 Roles and Responsibilities", level=2)
    doc.add_paragraph("Each role has defined responsibilities.")

    # Table in section 2.2 — tests table migration and border application
    roles_table = doc.add_table(rows=3, cols=2)
    roles_table.cell(0, 0).text = "Role"
    roles_table.cell(0, 1).text = "Responsibility"
    roles_table.cell(1, 0).text = "SE Lead"
    roles_table.cell(1, 1).text = "Oversees SE activities"
    roles_table.cell(2, 0).text = "CM Lead"
    roles_table.cell(2, 1).text = "Manages configuration baselines"

    doc.add_heading("3. SE Process", level=1)
    doc.add_paragraph("The SE process follows an iterative lifecycle model.")

    doc.add_heading("3.1 Planning", level=2)
    doc.add_paragraph("Planning activities occur at the start of each phase.")

    doc.add_heading("3.2 Old Configuration Process", level=2)
    doc.add_paragraph("This section describes the old config management approach.")

    doc.add_heading("3.3 Requirements Management", level=2)
    doc.add_paragraph("Requirements are tracked using a dedicated tool.")

    doc.save(output_path)
    print(f"  ✅ Source document created:      {output_path}")


def create_destination_template(output_path: str):
    """
    Create a minimal destination template (SEP v2) with:
      - Renamed sections to test fuzzy matching
      - Restructured subsections to test cross-parent mapping
      - Boilerplate placeholder text in each section
    """
    doc = Document()

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("[Boilerplate: Provide an introduction to the program.]")

    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph("[Boilerplate: Describe the purpose of this SEP.]")

    # Renamed from "1.2 Scope" — tests fuzzy matching
    doc.add_heading("1.2 Scope & Purpose", level=2)
    doc.add_paragraph("[Boilerplate: Describe the scope of this document.]")

    # Section 1.3 Background removed in v2 — source 1.3 will be unmapped

    doc.add_heading("2. Management Overview", level=1)
    doc.add_paragraph("[Boilerplate: Describe the management structure.]")

    doc.add_heading("2.1 Organizational Structure", level=2)
    doc.add_paragraph("[Boilerplate: Describe the organizational structure.]")

    doc.add_heading("2.2 Roles and Responsibilities", level=2)
    doc.add_paragraph("[Boilerplate: List roles and responsibilities.]")

    # Requirements Management moved from 3.3 to 2.3 — tests cross-parent mapping
    doc.add_heading("2.3 Requirements Management", level=2)
    doc.add_paragraph("[Boilerplate: Describe requirements management.]")

    doc.add_heading("3. SE Process", level=1)
    doc.add_paragraph("[Boilerplate: Describe the SE process.]")

    doc.add_heading("3.1 Planning", level=2)
    doc.add_paragraph("[Boilerplate: Describe planning activities.]")

    # Renamed from "3.2 Old Configuration Process"
    doc.add_heading("3.2 Configuration Management", level=2)
    doc.add_paragraph("[Boilerplate: Describe configuration management.]")

    doc.save(output_path)
    print(f"  ✅ Destination template created:  {output_path}")


# =============================================================================
# Individual test functions
# =============================================================================

def test_empty_section_filtering(source_sections: list) -> bool:
    """Verify that blank headings are not included in parsed sections."""
    blank_sections = [s for s in source_sections if not s.title.strip()]
    if blank_sections:
        print(f"  ❌ FAIL — {len(blank_sections)} blank section(s) not filtered out")
        return False
    print("  ✅ PASS — Empty section filter working correctly")
    return True


def test_auto_mapping(mapping_results: list) -> bool:
    """Verify basic auto-mapping statistics are reasonable."""
    auto_count     = sum(1 for r in mapping_results if r.status == MappingStatus.AUTO)
    review_count   = sum(1 for r in mapping_results if r.status == MappingStatus.REVIEW)
    unmapped_count = sum(1 for r in mapping_results if r.status == MappingStatus.UNMAPPED)

    print(f"  Mapping results: 🟢 {auto_count} auto  "
          f"🟡 {review_count} review  🔴 {unmapped_count} unmapped")

    # We expect at least some sections to be auto-mapped
    if auto_count == 0:
        print("  ❌ FAIL — No sections were auto-mapped")
        return False

    # original_status must be set on every result (not tacked on as raw attribute)
    for result in mapping_results:
        if not hasattr(result, 'original_status'):
            print("  ❌ FAIL — original_status field missing from MappingResult")
            return False

    print("  ✅ PASS — Auto-mapper produced results with correct structure")
    return True


def test_match_method_enum(mapping_results: list) -> bool:
    """Verify match_method uses MatchMethod enum values, not raw strings."""
    for result in mapping_results:
        if not isinstance(result.match_method, MatchMethod):
            print(f"  ❌ FAIL — match_method is '{type(result.match_method)}' "
                  f"not MatchMethod enum")
            return False
    print("  ✅ PASS — match_method uses MatchMethod enum correctly")
    return True


def test_migration(
    mapping_results: list,
    source_parser:   DocumentParser,
    dest_parser:     DocumentParser,
    dest_sections:   list,
    output_path:     str,
    mode:            MigrationMode,
) -> MigrationReport:
    """Run migration in the specified mode and return the report."""
    # Set all mapped results to the specified mode for this test pass
    for result in mapping_results:
        if result.dest_section is not None:
            result.migration_mode = mode

    # Re-parse for fresh Document objects
    fresh_source = DocumentParser(source_parser.filepath)
    fresh_source.parse()
    fresh_dest   = DocumentParser(dest_parser.filepath)
    fresh_dest.parse()

    migrator = ContentMigrator(
        mapping_results = mapping_results,
        source_document = fresh_source.document,
        dest_document   = fresh_dest.document,
        dest_sections   = dest_sections,
        output_path     = output_path,
    )
    report = migrator.run()

    mode_label = mode.value.capitalize()
    if report.errors:
        print(f"  ❌ {mode_label} mode — {len(report.errors)} error(s):")
        for error in report.errors:
            print(f"     {error}")
    else:
        print(f"  ✅ PASS — {mode_label} mode: "
              f"{report.total_migrated} migrated, "
              f"{len(report.skipped_sections)} skipped, "
              f"{len(report.unmapped_sections)} unmapped, "
              f"0 errors")

    return report


def test_boilerplate_detection(output_path: str, template_path: str) -> bool:
    """Run boilerplate detection and verify it completes without error."""
    try:
        detector    = BoilerplateDetector(template_path, output_path)
        match_count = detector.run()
        print(f"  ✅ PASS — Boilerplate detection: {match_count} match(es) recolored blue")
        return True
    except Exception as error:
        print(f"  ❌ FAIL — Boilerplate detection error: {error}")
        return False


def test_report_export(mapping_results: list, report_path: str) -> bool:
    """Verify the CSV report exports and contains the correct columns."""
    try:
        exporter = MigrationReportExporter(mapping_results)
        exporter.export(report_path)

        # Read back and verify columns
        import polars as pl
        df = pl.read_csv(report_path)
        expected_columns = {
            "source_section",
            "destination_section",
            "migration_mode",
            "outcome",
        }
        actual_columns = set(df.columns)
        if not expected_columns.issubset(actual_columns):
            missing = expected_columns - actual_columns
            print(f"  ❌ FAIL — Report missing columns: {missing}")
            return False

        print(f"  ✅ PASS — Report exported with correct columns "
              f"({len(df)} rows): {report_path}")
        return True
    except Exception as error:
        print(f"  ❌ FAIL — Report export error: {error}")
        return False


# =============================================================================
# Main test runner
# =============================================================================

def run_all_tests():
    print("\n" + "=" * 65)
    print("  Document Migration Tool v3 — Full Pipeline Test")
    print("=" * 65)

    test_dir = Path("/tmp/sep_test_v3")
    test_dir.mkdir(exist_ok=True)

    source_path       = str(test_dir / "source_sep_v1.docx")
    template_path     = str(test_dir / "template_sep_v2.docx")
    output_append     = str(test_dir / "output_append_mode.docx")
    output_replace    = str(test_dir / "output_replace_mode.docx")
    report_path       = str(test_dir / "migration_report.csv")

    all_passed = True

    # ── Create sample documents ───────────────────────────────────────────────
    print("\n[1] Creating sample documents…")
    create_source_document(source_path)
    create_destination_template(template_path)

    # ── Parse documents ───────────────────────────────────────────────────────
    print("\n[2] Parsing documents…")
    source_parser    = DocumentParser(source_path)
    source_sections  = source_parser.parse()
    dest_parser      = DocumentParser(template_path)
    dest_sections    = dest_parser.parse()
    print(f"  Source sections:      {len(source_sections)}")
    print(f"  Destination sections: {len(dest_sections)}")

    for section in source_sections:
        indent = "  " * (section.heading_level - 1)
        print(f"    {indent}[L{section.heading_level}] {section.title} "
              f"({len(section.content_blocks)} block(s))")

    # ── Test: empty section filtering ─────────────────────────────────────────
    print("\n[3] Testing empty section filter…")
    all_passed &= test_empty_section_filtering(source_sections)

    # ── Auto-map ──────────────────────────────────────────────────────────────
    print("\n[4] Running auto-mapper…")
    mapper          = AutoMapper(source_sections, dest_sections)
    mapping_results = mapper.run()

    for result in mapping_results:
        dest_title = result.dest_section.title if result.dest_section else "(none)"
        print(f"  {result.status_icon} {result.source_section.title!r:40s} "
              f"→ {dest_title!r}")

    # ── Test: auto-mapping correctness ────────────────────────────────────────
    print("\n[5] Testing auto-mapper output…")
    all_passed &= test_auto_mapping(mapping_results)

    # ── Test: MatchMethod enum ────────────────────────────────────────────────
    print("\n[6] Testing MatchMethod enum usage…")
    all_passed &= test_match_method_enum(mapping_results)

    # ── Test: Append mode migration ───────────────────────────────────────────
    print("\n[7] Testing migration — Append mode…")
    import copy as _copy
    append_results = _copy.deepcopy(mapping_results)
    append_report  = test_migration(
        append_results, source_parser, dest_parser, dest_sections,
        output_append, MigrationMode.APPEND,
    )
    all_passed &= (len(append_report.errors) == 0)

    # ── Test: Replace mode migration ──────────────────────────────────────────
    print("\n[8] Testing migration — Replace mode…")
    replace_results = _copy.deepcopy(mapping_results)
    replace_report  = test_migration(
        replace_results, source_parser, dest_parser, dest_sections,
        output_replace, MigrationMode.REPLACE,
    )
    all_passed &= (len(replace_report.errors) == 0)

    # ── Test: Boilerplate detection ───────────────────────────────────────────
    print("\n[9] Testing boilerplate detection…")
    all_passed &= test_boilerplate_detection(output_append, template_path)

    # ── Test: Report export ───────────────────────────────────────────────────
    print("\n[10] Testing migration report export…")
    all_passed &= test_report_export(mapping_results, report_path)

    # ── Verify output files exist ─────────────────────────────────────────────
    print("\n[11] Verifying output files exist…")
    for label, path in [
        ("Append mode output",  output_append),
        ("Replace mode output", output_replace),
        ("Migration report",    report_path),
    ]:
        if Path(path).exists():
            print(f"  ✅ {label}: {path}")
        else:
            print(f"  ❌ {label} NOT FOUND: {path}")
            all_passed = False

    # ── Final result ──────────────────────────────────────────────────────────
    print("\n" + "=" * 65)
    if all_passed:
        print("  ✅  ALL TESTS PASSED")
    else:
        print("  ❌  SOME TESTS FAILED — review output above")
    print("=" * 65 + "\n")

    return all_passed


if __name__ == "__main__":
    success = run_all_tests()
    sys.exit(0 if success else 1)
